"""
Generatore della "Dichiarazione sostitutiva dell'atto di notorieta'"
(Art. 47 D.P.R. 445/2000) per il monitoraggio e la tassazione delle
cripto-attivita', nel formato usato dai CAF.

Produce un file Word (.docx) con lo schema a punti: valori RW per asset,
acquisti dell'anno, vendite con prezzo medio di carico e plusvalenze, clausole
standard, luogo/data/firma.
"""

from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


@dataclass
class Anagrafica:
    nome: str
    nato_a: str
    nato_il: str
    codice_fiscale: str
    residente_a: str
    indirizzo: str
    caf: str = "Caf Veneto Srl - CAAF CGIL NORDEST"
    luogo: str = ""
    data: str = ""


@dataclass
class AssetRW:
    """Un asset da riportare nel Quadro RW."""
    denominazione: str
    wallet: str = ""
    valore_iniziale: float = 0.0
    valore_finale: float = 0.0
    # (data, quantita, controvalore) se acquisito durante l'anno
    acquisto: Optional[Tuple[str, float, float]] = None
    staking: bool = False
    solo_monitoraggio: bool = False   # imposta di bollo gia' autoliquidata


@dataclass
class VenditaAsset:
    denominazione: str
    vendite: List[Tuple[str, float, float]] = field(default_factory=list)   # (data, qta, controvalore)
    pmc: float = 0.0
    controvalore_carico: float = 0.0
    plusvalenze: List[Tuple[str, float]] = field(default_factory=list)      # (data, importo)


def _eur(x: float) -> str:
    """Formato italiano: 1.234,56."""
    return f"{x:,.2f}".translate(str.maketrans({",": ".", ".": ","}))


def _num(x: float) -> str:
    """Quantita' senza zeri finali inutili (14.21, 25, 0.068)."""
    s = f"{x:.8f}".rstrip("0").rstrip(".")
    return s.translate(str.maketrans({".": ","})) if "." in s else s


def dati_da_summary(
    summary: Dict,
    wallet_map: Optional[Dict[str, str]] = None,
    denom_map: Optional[Dict[str, str]] = None,
    staking: Optional[List[str]] = None,
    solo_monitoraggio: Optional[List[str]] = None,
) -> Tuple[List[AssetRW], List[VenditaAsset]]:
    """
    Costruisce (assets_rw, vendite) dal risultato di TaxCalculator.get_tax_summary().

    Args:
        summary: dict con "quadro_rw", "taxable_transactions", "all_transactions", "year".
        wallet_map: token -> nome del wallet/exchange ("wallet Binance", ...).
        denom_map: token -> denominazione estesa ("Bitcoin (BTC)"); default il token.
        staking / solo_monitoraggio: elenchi di token con quelle caratteristiche.
    """
    wallet_map = wallet_map or {}
    denom_map = denom_map or {}
    staking = set(staking or [])
    solo_monitoraggio = set(solo_monitoraggio or [])

    def denom(tok: str) -> str:
        return denom_map.get(tok, tok)

    anno = summary.get("year", "")

    # Acquisti valorizzati dell'anno, aggregati per token (per la frase "Di aver acquistato...").
    acq_anno: Dict[str, dict] = {}
    qta_acq_tot: Dict[str, float] = {}
    for tx in summary.get("all_transactions", []):
        if str(tx.get("type", "")).lower() != "buy" or not tx.get("price"):
            continue
        tok = tx["token"]
        d = acq_anno.setdefault(tok, {"n": 0, "prima_data": str(tx["date"]), "qta": 0.0, "costo": 0.0})
        d["n"] += 1
        d["qta"] += float(tx["amount"])
        d["costo"] += float(tx["total"])
        qta_acq_tot[tok] = qta_acq_tot.get(tok, 0.0) + float(tx["amount"])

    assets: List[AssetRW] = []
    for r in summary.get("quadro_rw", []):
        tok = r["token"]
        acq = None
        if r.get("quantita_iniziale", 0) <= 1e-9 and tok in acq_anno:
            d = acq_anno[tok]
            data_lbl = d["prima_data"] if d["n"] == 1 else f"nel corso del {anno}"
            acq = (data_lbl, d["qta"], d["costo"])
        assets.append(AssetRW(
            denominazione=denom(tok),
            wallet=wallet_map.get(tok, ""),
            valore_iniziale=float(r.get("valore_iniziale", 0.0)),
            valore_finale=float(r.get("valore_finale", 0.0)),
            acquisto=acq,
            staking=tok in staking,
            solo_monitoraggio=tok in solo_monitoraggio,
        ))

    # Vendite: raggruppa per token le cessioni dell'anno e le plusvalenze LIFO/PMC.
    vendite_per_tok: Dict[str, List[Tuple[str, float, float]]] = {}
    for tx in summary.get("all_transactions", []):
        if str(tx.get("type", "")).lower() == "sell" and tx.get("price"):
            vendite_per_tok.setdefault(tx["token"], []).append(
                (str(tx["date"]), float(tx["amount"]), float(tx["total"]))
            )

    plus_per_tok: Dict[str, List[Tuple[str, float]]] = {}
    pmc_per_tok: Dict[str, float] = {}
    for e in summary.get("taxable_transactions", []):
        plus_per_tok.setdefault(e["token"], []).append((str(e["date"]), float(e["gain"])))
        pmc_per_tok[e["token"]] = float(e.get("buy_price", 0.0))

    vendite: List[VenditaAsset] = []
    for tok, vs in vendite_per_tok.items():
        pmc = pmc_per_tok.get(tok, 0.0)
        vendite.append(VenditaAsset(
            denominazione=denom(tok),
            vendite=vs,
            pmc=pmc,
            controvalore_carico=pmc * qta_acq_tot.get(tok, 0.0),
            plusvalenze=plus_per_tok.get(tok, []),
        ))
    return assets, vendite


def genera_dichiarazione_docx(
    anagrafica: Anagrafica,
    assets_rw: List[AssetRW],
    vendite: List[VenditaAsset],
    anno: int,
    output_path: str,
    note_extra: Optional[List[str]] = None,
) -> bool:
    """Scrive la dichiarazione in `output_path` (.docx). True se riesce."""
    try:
        doc = Document()
        for s in doc.sections:
            s.left_margin = s.right_margin = Pt(56)

        def titolo(t):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(t)
            run.bold = True
            return p

        def par(t="", bold=False):
            p = doc.add_paragraph()
            if t:
                run = p.add_run(t)
                run.bold = bold
            return p

        def punto(t):
            doc.add_paragraph(t, style="List Bullet")

        titolo("DICHIARAZIONE SOSTITUTIVA DELL'ATTO DI NOTORIETA'")
        titolo("(Art. 47 D.P.R. 28 dicembre 2000, n. 445)")
        titolo("DA PRESENTARE ALLA PUBBLICA AMMINISTRAZIONE O AI GESTORI DI PUBBLICI SERVIZI.")
        par()

        a = anagrafica
        par(f"Il sottoscritto {a.nome} nato a {a.nato_a} il {a.nato_il} e residente a "
            f"{a.residente_a} in {a.indirizzo}")
        par(f"Codice Fiscale {a.codice_fiscale}")
        par()
        par("consapevole delle sanzioni penali, nel caso di dichiarazioni non veritiere, di "
            "formazione o uso di atti falsi, richiamate dall'art. 76 del D.P.R. 445 del 28 dicembre 2000")
        titolo("DICHIARA ED ATTESTA")
        par()

        for asset in assets_rw:
            if asset.acquisto:
                data_a, qta_a, ctv_a = asset.acquisto
                quando = f"in data {data_a}" if "/" in str(data_a) else f"{data_a}"
                unita = "unita'" if "/" in str(data_a) else "unita' complessive"
                punto(
                    f"Di aver acquistato criptoattivita' denominata {asset.denominazione} {quando} "
                    f"per {_num(qta_a)} {unita} al controvalore di {_eur(ctv_a)} EUR "
                    f"mentre il controvalore al 31/12/{anno} era pari a {_eur(asset.valore_finale)} EUR;"
                )
                continue

            frase = (
                f"Che il valore della Criptovaluta {asset.denominazione}"
                f"{' detenuta presso ' + asset.wallet if asset.wallet else ''}"
                f"{' soggetta a vincolo di disponibilità (staking)' if asset.staking else ''}"
                f" era pari a {_eur(asset.valore_iniziale)} EUR alla data del 01/01/{anno} "
                f"mentre il valore era pari a {_eur(asset.valore_finale)} EUR alla data del 31/12/{anno};"
            )
            if asset.solo_monitoraggio:
                frase += (
                    " Su tale importo e' stata calcolata e versata da me in autoliquidazione "
                    "l'imposta di bollo sulle criptoattivita', e quindi intendo indicare tale "
                    "criptoattivita' solo ai fini del monitoraggio;"
                )
            punto(frase)

        for v in vendite:
            elenco = " e ".join(
                f"{_num(q)} unita' in data {d} per un controvalore di {_eur(c)} EUR"
                for d, q, c in v.vendite
            )
            punto(
                f"Di aver effettuato le seguenti vendite di criptovaluta {v.denominazione}: {elenco}. "
                f"Tale criptoattivita' era detenuta ad un prezzo medio di carico di {_eur(v.pmc)} EUR "
                f"per un controvalore di {_eur(v.controvalore_carico)} EUR."
            )
            if v.plusvalenze:
                plus = " e ".join(f"pari a {_eur(g)} EUR in data {d}" for d, g in v.plusvalenze)
                coda = ", da indicare al lordo della franchigia di 2000 EUR;" if anno <= 2024 \
                    else ", da indicare nel quadro RT (Sezione V) della dichiarazione dei redditi;"
                punto(f"Che le suddette vendite hanno determinato una plusvalenza imponibile {plus}{coda}")

        for n in (note_extra or []):
            punto(n)

        punto("Che in caso di accertamento da parte dell'agenzia delle entrate e' in grado di "
              "produrre documentazione che comprova quanto dichiarato con la presente autodichiarazione;")
        punto("Di non possedere altri patrimoni esteri, sia di natura immobiliare che mobiliare o "
              "finanziaria, in qualunque nazione o stato estero detenuti;")
        punto("Che e' stato informato che l'infedele presentazione della dichiarazione dei redditi da "
              "parte del contribuente comporta l'erogazione di sanzioni ed interessi oltre il pagamento "
              "delle imposte dovute a seguito della liquidazione della posizione fiscale da parte "
              "dell'Agenzia delle Entrate;")
        punto(f"Di esonerare il {a.caf} da qualsiasi responsabilita' presente e futura derivante "
              "dall'infedele presentazione della dichiarazione dei redditi e di non aver nulla a "
              "pretendere dai suddetti a titolo di imposta dovuta, sanzioni ed interessi da essa derivanti.")
        par()
        par("Dichiara, altresi', di essere informato, ai sensi e per gli effetti di cui al D. Lgs. "
            "N. 196/2003, che i dati personali raccolti saranno trattati, anche con strumenti "
            "informatici, esclusivamente nell'ambito del procedimento per il quale la presente "
            "dichiarazione viene resa.")
        par()
        par(f"{a.luogo}, {a.data}")
        par()
        p = par("Il Dichiarante", bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2 = par("......................................................")
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        par()
        p3 = par("Ai sensi dell'art. 38, D.P.R. 445 del 28 dicembre 2000, la dichiarazione e' "
                 "sottoscritta dall'interessato in presenza del dipendente addetto ovvero "
                 "sottoscritta e inviata insieme alla fotocopia, non autenticata di un documento "
                 "di identita' del dichiarante.")
        for r in p3.runs:
            r.font.size = Pt(8)

        doc.save(output_path)
        return True
    except Exception as e:  # pragma: no cover
        print(f"Errore generazione dichiarazione: {e}")
        return False
